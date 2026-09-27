"""
Local text extraction from resume files (spec §10). PDF via PyMuPDF, DOCX
via python-docx — no binary resume content goes to an external LLM when
local extraction is sufficient (spec §10: "Do not send a complete binary
resume to an external LLM when locally extracted text is sufficient").

OCR fallback (pytesseract + pdf2image, both local) only triggers when
extraction quality looks too poor to be a real text layer — i.e. the PDF is
probably scanned images, not selectable text.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

MIN_CHARS_PER_PAGE_BEFORE_OCR = 40  # heuristic threshold, spec §10 "assess extraction quality"


@dataclass
class ExtractedText:
    text: str
    page_count: int
    used_ocr: bool


def extract_pdf_text(data: bytes) -> ExtractedText:
    import fitz  # PyMuPDF

    doc = fitz.open(stream=data, filetype="pdf")
    try:
        pages_text = [page.get_text() for page in doc]
        page_count = doc.page_count
    finally:
        doc.close()

    text = "\n\n".join(pages_text)

    if _extraction_quality_is_poor(text, page_count):
        ocr_text = _ocr_pdf(data)
        if ocr_text.strip():
            return ExtractedText(text=ocr_text, page_count=page_count, used_ocr=True)

    return ExtractedText(text=text, page_count=page_count, used_ocr=False)


def extract_docx_text(data: bytes) -> ExtractedText:
    import docx

    document = docx.Document(io.BytesIO(data))

    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    # Resumes sometimes use tables for layout (skills grids, two-column
    # contact blocks) — walk them too, or that content silently disappears.
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    text = "\n".join(parts)
    return ExtractedText(text=text, page_count=1, used_ocr=False)  # DOCX has no fixed page count pre-render


def _extraction_quality_is_poor(text: str, page_count: int) -> bool:
    if page_count == 0:
        return True
    avg_chars_per_page = len(text.strip()) / page_count
    return avg_chars_per_page < MIN_CHARS_PER_PAGE_BEFORE_OCR


def _ocr_pdf(data: bytes) -> str:
    """
    Local OCR fallback (spec §10). Requires poppler (for pdf2image) to be
    installed on the host — this is a system dependency, not a pip package,
    and is documented in backend/README.md's OCR section.
    """
    import pytesseract
    from pdf2image import convert_from_bytes

    images = convert_from_bytes(data)
    pages = [pytesseract.image_to_string(image) for image in images]
    return "\n\n".join(pages)
