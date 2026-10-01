"""
Deterministic rendering (spec §27): the same content always produces the same
document. Single column, plain text, no images or text boxes, so ATS parsers
read it cleanly. Employer names, titles, dates, education and certifications
come from the candidate's stored profile, never from the model.
"""

from __future__ import annotations

import datetime as dt
import hashlib
from dataclasses import dataclass, field
from html import escape
from pathlib import Path

from app.models.candidate import CandidateProfile
from app.services.llm.schemas import TailoredResumeContent

TEMPLATE_VERSION = "1"


@dataclass
class RenderJob:
    heading: str
    dates: str
    bullets: list[str]


@dataclass
class RenderableResume:
    name: str
    contact_line: str
    summary: str
    jobs: list[RenderJob] = field(default_factory=list)
    education: list[str] = field(default_factory=list)
    certifications: list[str] = field(default_factory=list)
    skills: list[str] = field(default_factory=list)


def _fmt_date(d: dt.date | None) -> str:
    return d.strftime("%b %Y") if d else ""


def build_renderable(profile: CandidateProfile, content: TailoredResumeContent) -> RenderableResume:
    by_id = {e.id: e for e in profile.employment_history}
    jobs: list[RenderJob] = []
    for exp in content.experience:
        entry = by_id.get(exp.employment_id)
        if entry is None or not exp.bullets:
            continue
        dates = f"{_fmt_date(entry.start_date) or '?'} - {_fmt_date(entry.end_date) or 'Present'}"
        jobs.append(
            RenderJob(
                heading=f"{entry.title}, {entry.employer}",
                dates=dates,
                bullets=[b.text for b in exp.bullets],
            )
        )

    education = [
        " ".join(p for p in [e.degree, e.field] if p).strip() + (", " if (e.degree or e.field) else "") + e.institution
        for e in profile.education
    ]
    certs = [c.certification + (f" ({c.issuer})" if c.issuer else "") for c in profile.certifications]
    contact = " | ".join(p for p in [profile.email, profile.phone, profile.location] if p)

    return RenderableResume(
        name=profile.name,
        contact_line=contact,
        summary=content.summary,
        jobs=jobs,
        education=education,
        certifications=certs,
        skills=content.skills,
    )


def content_hash(resume: RenderableResume) -> str:
    return hashlib.sha256(repr(resume).encode("utf-8")).hexdigest()


def render_docx(resume: RenderableResume, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading(resume.name, level=0)
    if resume.contact_line:
        doc.add_paragraph(resume.contact_line)
    if resume.summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(resume.summary)
    if resume.jobs:
        doc.add_heading("Experience", level=1)
        for job in resume.jobs:
            p = doc.add_paragraph()
            p.add_run(job.heading).bold = True
            p.add_run(f"  {job.dates}")
            for bullet in job.bullets:
                doc.add_paragraph(bullet, style="List Bullet")
    if resume.education:
        doc.add_heading("Education", level=1)
        for line in resume.education:
            doc.add_paragraph(line)
    if resume.certifications:
        doc.add_heading("Certifications", level=1)
        for line in resume.certifications:
            doc.add_paragraph(line)
    if resume.skills:
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(resume.skills))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


_CSS = """
body { font-family: sans-serif; font-size: 10pt; }
h1 { font-size: 18pt; margin-bottom: 2pt; }
h2 { font-size: 11pt; margin-top: 10pt; margin-bottom: 3pt; border-bottom: 1px solid #000; }
p { margin: 0 0 3pt 0; }
li { margin-bottom: 2pt; }
.job { font-weight: bold; }
"""


def _resume_html(resume: RenderableResume) -> str:
    e = escape
    parts = [f"<h1>{e(resume.name)}</h1>"]
    if resume.contact_line:
        parts.append(f"<p>{e(resume.contact_line)}</p>")
    if resume.summary:
        parts.append(f"<h2>Summary</h2><p>{e(resume.summary)}</p>")
    if resume.jobs:
        parts.append("<h2>Experience</h2>")
        for job in resume.jobs:
            parts.append(f'<p class="job">{e(job.heading)} &#8211; {e(job.dates)}</p><ul>')
            parts.extend(f"<li>{e(b)}</li>" for b in job.bullets)
            parts.append("</ul>")
    if resume.education:
        parts.append("<h2>Education</h2>" + "".join(f"<p>{e(x)}</p>" for x in resume.education))
    if resume.certifications:
        parts.append("<h2>Certifications</h2>" + "".join(f"<p>{e(x)}</p>" for x in resume.certifications))
    if resume.skills:
        parts.append(f"<h2>Skills</h2><p>{e(', '.join(resume.skills))}</p>")
    return "".join(parts)


def render_pdf(resume: RenderableResume, path: Path) -> None:
    # PyMuPDF is already a dependency (resume parsing); its Story API lays out
    # simple HTML across pages, so no extra PDF library is needed.
    import fitz

    path.parent.mkdir(parents=True, exist_ok=True)
    story = fitz.Story(html=_resume_html(resume), user_css=_CSS)
    writer = fitz.DocumentWriter(str(path))
    page_rect = fitz.paper_rect("letter")
    content_rect = page_rect + (54, 54, -54, -54)
    more = True
    while more:
        device = writer.begin_page(page_rect)
        more, _ = story.place(content_rect)
        story.draw(device)
        writer.end_page()
    writer.close()
