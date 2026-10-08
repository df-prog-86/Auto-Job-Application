"""Extension fill flow (context + report) and the Needs Attention list."""

from __future__ import annotations

import datetime as dt

_COMMIT = {
    "extraction": {
        "contact": {"name": "Jane Q Doe", "email": "jane@example.com", "phone": "555-0100", "location": "Boston, MA"},
        "employment": [],
        "education": [],
        "skills": [],
        "certifications": [],
        "projects": [],
    },
    "approved_claims": [],
    "resume_filename": "resume.docx",
}


def _token(client) -> dict:
    secret = client.post("/api/v1/system/pairing-secret").json()["pairing_secret"]
    token = client.post(
        "/api/v1/system/pair", json={"pairing_secret": secret, "extension_origin": "chrome-extension://t"}
    ).json()["extension_token"]
    return {"X-Extension-Token": token}


def _job(
    SessionLocal, *, url="https://job-boards.greenhouse.io/acme/jobs/123", proceeding=True, resume=True, resume_path="/tmp/Jane Doe_Resume_Acme_2026.docx"
) -> int:
    from app.models.documents import GeneratedDocument
    from app.models.jobs import Job

    now = dt.datetime.now(dt.timezone.utc)
    with SessionLocal() as db:
        job = Job(
            canonical_job_key=f"k-{url}",
            company="Acme",
            normalized_company="acme",
            title="Analyst",
            normalized_title="analyst",
            canonical_application_url=url,
            first_seen=now,
            last_seen=now,
            application_status="proceeding" if proceeding else "not_started",
        )
        db.add(job)
        db.flush()
        if resume:
            db.add(
                GeneratedDocument(
                    job_id=job.id,
                    document_type="resume",
                    local_path=resume_path,
                    format="docx",
                    generated_at=now,
                    template_version="t",
                    content_hash="h",
                )
            )
        db.commit()
        return job.id


def test_apply_endpoints_require_the_extension_token(app_and_db):
    client, _ = app_and_db
    assert client.post("/api/v1/apply/context", json={"url": "https://x.com"}).status_code == 401
    assert client.post("/api/v1/apply/report", json={"job_id": 1}).status_code == 401


def test_context_needs_a_profile(app_and_db):
    client, _ = app_and_db
    out = client.post("/api/v1/apply/context", json={"url": "https://x.com"}, headers=_token(client)).json()
    assert out["candidate"] is None and "Profile" in out["problem"]


def test_context_matches_job_and_returns_facts_answers_and_resume(app_and_db):
    client, SessionLocal = app_and_db
    client.post("/api/v1/profile/commit", json=_COMMIT)
    client.put("/api/v1/profile/answers/work_authorization", json={"value_type": "str", "value": "US citizen"})
    client.put("/api/v1/profile/answers/sponsorship_required", json={"value_type": "bool", "value": False})
    client.put("/api/v1/profile/answers/phone_country", json={"value_type": "str", "value": "United States (+1)"})
    client.put("/api/v1/profile/answers/eeo_veteran", json={"value_type": "str", "value": "not_protected"})
    client.put("/api/v1/profile/answers/phone_device_type", json={"value_type": "str", "value": "Mobile"})
    client.put("/api/v1/profile/answers/address_line1", json={"value_type": "str", "value": "1 Main St"})
    client.put("/api/v1/profile/answers/postal_code", json={"value_type": "str", "value": "02118"})
    client.patch("/api/v1/profile", json={"preferred_name": "Janie"})
    client.put("/api/v1/profile/answers/q:why do you want this role", json={"value_type": "str", "value": "Mission fit."})
    job_id = _job(SessionLocal)
    headers = _token(client)

    out = client.post(
        "/api/v1/apply/context",
        json={"url": "https://job-boards.greenhouse.io/acme/jobs/123?gh_src=abc#app"},
        headers=headers,
    ).json()
    assert out["problem"] is None
    assert out["job"]["id"] == job_id
    assert out["candidate"]["first_name"] == "Jane"
    assert out["candidate"]["last_name"] == "Q Doe"
    assert out["candidate"]["email"] == "jane@example.com"
    assert out["candidate"]["preferred_name"] == "Janie"
    assert out["answers"] == {
        "work_authorization": "US citizen",
        "sponsorship_required": False,
        "phone_country": "United States (+1)",
        "eeo_veteran": "not_protected",
        "phone_device_type": "Mobile",
        "address_line1": "1 Main St",
        "postal_code": "02118",
    }
    assert out["learned_answers"] == {"why do you want this role": "Mission fit."}
    assert out["resume"]["filename"] == "Jane Doe_Resume_Acme_2026.docx"


def test_context_stops_before_proceed_or_resume(app_and_db):
    client, SessionLocal = app_and_db
    client.post("/api/v1/profile/commit", json=_COMMIT)
    headers = _token(client)
    _job(SessionLocal, url="https://a.com/one", proceeding=False)
    _job(SessionLocal, url="https://a.com/two", resume=False)

    not_proceeding = client.post("/api/v1/apply/context", json={"url": "https://a.com/one"}, headers=headers).json()
    assert "Proceed" in not_proceeding["problem"] and not_proceeding["candidate"] is None

    no_resume = client.post("/api/v1/apply/context", json={"url": "https://a.com/two"}, headers=headers).json()
    assert "Choose a resume" in no_resume["problem"] and no_resume["candidate"] is None


def test_context_lists_choices_when_page_is_unknown(app_and_db):
    client, SessionLocal = app_and_db
    client.post("/api/v1/profile/commit", json=_COMMIT)
    _job(SessionLocal, url="https://a.com/one")
    out = client.post("/api/v1/apply/context", json={"url": "https://elsewhere.com/x"}, headers=_token(client)).json()
    assert out["job"] is None and len(out["candidates"]) == 1 and "pick" in out["problem"].lower()


def test_report_creates_questions_once_and_needs_attention_flow(app_and_db):
    client, SessionLocal = app_and_db
    client.post("/api/v1/profile/commit", json=_COMMIT)
    job_id = _job(SessionLocal)
    headers = _token(client)
    body = {
        "job_id": job_id,
        "page_url": "https://job-boards.greenhouse.io/acme/jobs/123",
        "filled_count": 5,
        "flagged": [
            {"label": "Why do you want to work here? *", "field_type": "textarea", "required": True},
            {"label": "How did you hear about us?", "field_type": "select", "options": ["Referral", "Other"]},
        ],
    }
    assert client.post("/api/v1/apply/report", json=body, headers=headers).status_code == 204
    assert client.post("/api/v1/apply/report", json=body, headers=headers).status_code == 204  # re-run: no duplicates

    open_items = client.get("/api/v1/needs-attention").json()
    assert len(open_items) == 2
    why = next(i for i in open_items if i["label"].startswith("Why do you"))
    assert why["company"] == "Acme" and why["required"] is True

    answered = client.post(f"/api/v1/needs-attention/{why['id']}/answer", json={"answer": "Mission fit."})
    assert answered.status_code == 200 and answered.json()["status"] == "answered"
    assert len(client.get("/api/v1/needs-attention").json()) == 1

    # The answer is remembered and handed back on the next fill, for any job.
    ctx = client.post(
        "/api/v1/apply/context", json={"url": "https://job-boards.greenhouse.io/acme/jobs/123"}, headers=headers
    ).json()
    assert ctx["learned_answers"] == {"why do you want to work here": "Mission fit."}

    # Answered questions don't come back when Fill runs again.
    client.post("/api/v1/apply/report", json=body, headers=headers)
    assert len(client.get("/api/v1/needs-attention").json()) == 1

    other = client.get("/api/v1/needs-attention").json()[0]
    assert client.post(f"/api/v1/needs-attention/{other['id']}/dismiss").json()["status"] == "dismissed"
    assert client.get("/api/v1/needs-attention").json() == []
    assert client.post(f"/api/v1/needs-attention/{why['id']}/answer", json={"answer": "  "}).status_code == 422


def test_deleting_a_job_removes_its_questions(app_and_db):
    client, SessionLocal = app_and_db
    client.post("/api/v1/profile/commit", json=_COMMIT)
    job_id = _job(SessionLocal)
    client.post(
        "/api/v1/apply/report",
        json={"job_id": job_id, "flagged": [{"label": "Anything else?"}]},
        headers=_token(client),
    )
    assert len(client.get("/api/v1/needs-attention").json()) == 1
    assert client.delete(f"/api/v1/jobs/{job_id}").status_code == 204
    assert client.get("/api/v1/needs-attention").json() == []


def test_ashby_posting_and_application_urls_match_and_recent_role_is_returned(app_and_db):
    from app.services.apply.questions import urls_match

    posting = "https://jobs.ashbyhq.com/chartis/5abc9d3f-9ade-4a3e-a8c6-0d76c98b3fa4"
    assert urls_match(posting + "/application?utm=x", posting)
    assert not urls_match("https://jobs.ashbyhq.com/chartis/other-id/application", posting)

    client, SessionLocal = app_and_db
    client.post("/api/v1/profile/commit", json=_COMMIT)
    client.post("/api/v1/profile/employment", json={"employer": "Old Co", "title": "Analyst", "start_date": "2015-01", "end_date": "2018-06"})
    client.post("/api/v1/profile/employment", json={"employer": "Huron", "title": "Manager", "start_date": "2019-03"})
    _job(SessionLocal, url=posting)
    out = client.post("/api/v1/apply/context", json={"url": posting + "/application"}, headers=_token(client)).json()
    assert out["problem"] is None
    assert out["candidate"]["recent_title"] == "Manager"
    assert out["candidate"]["recent_employer"] == "Huron"


def test_context_returns_roles_with_bullets_from_the_resume_used_for_the_job(app_and_db, monkeypatch, tmp_path):
    from docx import Document

    from app.config import settings

    monkeypatch.setattr(settings, "GENERATED_DOCUMENTS_DIR", str(tmp_path / "generated_documents"))
    client, SessionLocal = app_and_db
    commit = {
        **_COMMIT,
        "extraction": {
            **_COMMIT["extraction"],
            "employment": [
                {"employer": "Acme Corp", "title": "Senior Data Analyst", "start_date": "2021-04", "end_date": None, "source_text": "x"},
                {"employer": "Beta LLC", "title": "Analyst", "start_date": "2018-01", "end_date": "2021-03", "source_text": "y"},
            ],
            "education": [{"institution": "State University", "degree": "BS", "field": "Finance", "start_date": "2014-09", "end_date": "2018-05"}],
        },
    }
    client.post("/api/v1/profile/commit", json=commit)

    doc = Document()
    doc.add_paragraph("Senior Data Analyst, Acme Corp | 2021 - Present")
    doc.add_paragraph("Tailored bullet one", style="List Bullet")
    doc.add_paragraph("Tailored bullet two", style="List Bullet")
    doc.add_paragraph("Analyst, Beta LLC | 2018 - 2021")
    doc.add_paragraph("Older bullet", style="List Bullet")
    folder = tmp_path / "generated_documents" / "job_1"
    folder.mkdir(parents=True)
    path = folder / "resume.docx"
    doc.save(str(path))

    job_id = _job(SessionLocal, resume_path=str(path))
    out = client.post("/api/v1/apply/context", json={"url": "https://x.com", "job_id": job_id}, headers=_token(client)).json()
    assert [r["employer"] for r in out["experience"]] == ["Acme Corp", "Beta LLC"]  # current role first
    acme, beta = out["experience"]
    assert acme["current"] is True and acme["end_date"] is None and acme["start_date"] == "2021-04"
    assert acme["description"] == "\u2022 Tailored bullet one\n\u2022 Tailored bullet two"
    assert beta["current"] is False and beta["end_date"] == "2021-03" and beta["description"] == "\u2022 Older bullet"
    assert out["education"][0]["institution"] == "State University" and out["education"][0]["end_date"] == "2018-05"


def test_context_skills_put_the_ones_in_the_posting_first(app_and_db):
    client, SessionLocal = app_and_db
    client.post("/api/v1/profile/commit", json=_COMMIT)
    for name in ["Python", "Tableau", "SQL", "Excel", "sql"]:
        client.post("/api/v1/profile/skills", json={"canonical_skill": name})
    job_id = _job(SessionLocal)
    from app.models.jobs import Job

    with SessionLocal() as db:
        db.get(Job, job_id).description = "We need strong SQL and Tableau; Excelling is not a skill."
        db.commit()
    out = client.post("/api/v1/apply/context", json={"url": "https://x.com", "job_id": job_id}, headers=_token(client)).json()
    assert out["skills"] == ["Tableau", "SQL", "Python", "Excel"]  # posting words first, no duplicates, no partial-word hits


def test_education_gpa_and_certifications_are_editable_and_reach_the_apply_context(app_and_db):
    client, SessionLocal = app_and_db
    client.post("/api/v1/profile/commit", json=_COMMIT)

    profile = client.post("/api/v1/profile/education", json={"institution": "State University", "degree": "Bachelors Degree"}).json()
    edu_id = profile["education"][0]["id"]
    out = client.patch(f"/api/v1/profile/education/{edu_id}", json={"start_date": "2014-09", "end_date": "2018-05", "gpa": " 3.7 "}).json()
    assert out["education"][0]["gpa"] == "3.7" and out["education"][0]["start_date"] == "2014-09-01"

    created = client.post("/api/v1/profile/certifications", json={"certification": "Epic Clarity", "issuer": "Epic", "date": "2022-03-15"}).json()
    cert_id = created["certifications"][0]["id"]
    out = client.patch(f"/api/v1/profile/certifications/{cert_id}", json={"expiration": "2026-03-15", "certification": "  "}).json()
    assert out["certifications"][0]["certification"] == "Epic Clarity"  # a required name cannot be blanked
    assert out["certifications"][0]["expiration"] == "2026-03-15"

    job_id = _job(SessionLocal)
    ctx = client.post("/api/v1/apply/context", json={"url": "https://x.com", "job_id": job_id}, headers=_token(client)).json()
    assert ctx["education"][0]["gpa"] == "3.7" and ctx["education"][0]["start_date"] == "2014-09"
    assert ctx["certifications"] == [{"name": "Epic Clarity", "issuer": "Epic", "issued": "2022-03-15", "expires": "2026-03-15"}]

    assert client.delete(f"/api/v1/profile/certifications/{cert_id}").json()["certifications"] == []
    assert client.delete("/api/v1/profile/certifications/9999").status_code == 404


def test_answers_the_person_gives_on_a_form_are_remembered_and_can_be_forgotten(app_and_db):
    client, SessionLocal = app_and_db
    client.post("/api/v1/profile/commit", json=_COMMIT)
    job_id = _job(SessionLocal)
    headers = _token(client)
    client.post(
        "/api/v1/apply/report",
        json={"job_id": job_id, "flagged": [{"label": "Do you have a relative here? *", "field_type": "dropdown", "options": ["Yes", "No"], "required": True}]},
        headers=headers,
    )
    assert len(client.get("/api/v1/needs-attention").json()) == 1

    out = client.post(
        "/api/v1/apply/learned",
        json={
            "answers": [
                {"label": "Do you have a relative here?", "value": "No"},
                {"label": "I agree to the terms", "value": "Yes"},  # agreements are never remembered
                {"label": "Have you ever been fired?", "value": "No"},  # history questions are never remembered
                {"label": "Your password", "value": "x"},
            ]
        },
        headers=headers,
    ).json()
    assert out == {"saved": 1}
    assert client.get("/api/v1/needs-attention").json() == []  # settled by the saved answer

    ctx = client.post("/api/v1/apply/context", json={"url": "https://x.com", "job_id": job_id}, headers=headers).json()
    assert ctx["learned_answers"] == {"do you have a relative here": "No"}

    key = "q:do you have a relative here"
    shown = [a for a in client.get("/api/v1/profile/answers").json() if a["answer_key"] == key][0]
    assert shown["value"] == "No" and shown["explanatory_text"] == "Do you have a relative here?"
    assert client.delete("/api/v1/profile/answers/work_authorization").status_code == 400  # standard answers are edited, not deleted
    assert client.delete(f"/api/v1/profile/answers/{key}").status_code == 204
    assert [a for a in client.get("/api/v1/profile/answers").json() if a["answer_key"] == key] == []


def test_needs_attention_can_be_cleared_and_closes_what_the_profile_now_answers(app_and_db):
    client, SessionLocal = app_and_db
    client.post("/api/v1/profile/commit", json=_COMMIT)
    job_id = _job(SessionLocal)
    headers = _token(client)
    flagged = [
        {"label": "Address Line 1", "field_type": "text", "options": [], "required": True},
        {"label": "Do you have a relative here?", "field_type": "dropdown", "options": ["Yes", "No"], "required": True},
        {"label": "Preferred pronouns?", "field_type": "text", "options": [], "required": True},
    ]
    client.post("/api/v1/apply/report", json={"job_id": job_id, "flagged": flagged}, headers=headers)
    assert len(client.get("/api/v1/needs-attention").json()) == 3

    # Saving the street address on the Profile settles the matching question by itself.
    client.put("/api/v1/profile/answers/address_line1", json={"value_type": "str", "value": "1 Main St"})
    open_items = client.get("/api/v1/needs-attention").json()
    assert sorted(q["label"] for q in open_items) == ["Do you have a relative here?", "Preferred pronouns?"]

    # Clear some by id, then the rest.
    first = open_items[0]["id"]
    assert client.post("/api/v1/needs-attention/clear", json={"ids": [first]}).json() == {"cleared": 1}
    assert len(client.get("/api/v1/needs-attention").json()) == 1
    assert client.post("/api/v1/needs-attention/clear", json={}).json() == {"cleared": 1}
    assert client.get("/api/v1/needs-attention").json() == []


def test_extension_reports_a_submission_once_and_the_first_date_stands(app_and_db):
    client, SessionLocal = app_and_db
    job_id = _job(SessionLocal)
    headers = _token(client)
    assert client.post("/api/v1/apply/submitted", json={"job_id": job_id}).status_code == 401
    assert client.post("/api/v1/apply/submitted", json={"job_id": 9999}, headers=headers).status_code == 404
    assert client.post("/api/v1/apply/submitted", json={"job_id": job_id}, headers=headers).status_code == 204
    first = client.get(f"/api/v1/jobs/{job_id}").json()
    assert first["applied_via"] == "extension" and first["applied_at"]
    assert client.post("/api/v1/apply/submitted", json={"job_id": job_id}, headers=headers).status_code == 204
    assert client.get(f"/api/v1/jobs/{job_id}").json()["applied_at"] == first["applied_at"]


def test_a_fill_report_marks_the_application_as_started_once(app_and_db):
    client, SessionLocal = app_and_db
    job_id = _job(SessionLocal)
    headers = _token(client)
    assert client.get(f"/api/v1/jobs/{job_id}").json()["application_started_at"] is None
    body = {"job_id": job_id, "page_url": "https://job-boards.greenhouse.io/acme/jobs/123", "filled_count": 3, "flagged": []}
    assert client.post("/api/v1/apply/report", json=body, headers=headers).status_code == 204
    first = client.get(f"/api/v1/jobs/{job_id}").json()["application_started_at"]
    assert first
    client.post("/api/v1/apply/report", json=body, headers=headers)
    assert client.get(f"/api/v1/jobs/{job_id}").json()["application_started_at"] == first
