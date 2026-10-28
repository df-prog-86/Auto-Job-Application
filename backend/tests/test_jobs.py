"""
Job intake (paste-a-URL), scoring, the Proceed gate, delete/undo, and resume
tailoring, end to end through the API. No test touches the real network or a
real LLM: page fetching and the model router are replaced with fakes, and
every test uses its own temporary database and folders.
"""

from __future__ import annotations

import datetime as dt
import io
import json
from urllib.parse import unquote

import pytest
from docx import Document

from app.config import settings
from app.models.documents import GeneratedDocument
from app.services.discovery import manual_extraction
from app.services.discovery.manual_extraction import FetchedPage, ManualExtractionError
from app.services.llm.schemas import (
    JobPostingExtraction,
    PlannedBlock,
    PlannedBullet,
    ResumeJobMatchResult,
    TailorPlan,
)


_FULL_DESCRIPTION = (
    "<p>Analyze data with SQL. 3+ years of experience required.</p>"
    "<p>You will build reports, work with finance and operations teams, and explain results to leaders. "
    "Strong Excel and SQL skills are required, and experience with dashboards is a plus. "
    "The role is based in Boston with some flexibility to work from home during the week.</p>"
)


def _jsonld(title="Data Analyst", company="Globex", url="https://example.com/jobs/1", description=_FULL_DESCRIPTION) -> str:
    return json.dumps(
        {
            "@context": "https://schema.org",
            "@type": "JobPosting",
            "title": title,
            "hiringOrganization": {"@type": "Organization", "name": company},
            "description": description,
            "jobLocation": {"address": {"addressLocality": "Boston", "addressRegion": "MA"}},
            "url": url,
        }
    )


def _add_job(client, monkeypatch, **kwargs) -> dict:
    """Adds a job by link with the page fetch faked; returns the response JSON."""
    block = _jsonld(**kwargs)

    async def fake_fetch(url: str) -> FetchedPage:
        return FetchedPage(url=url, json_ld_blocks=[block], body_text="", title="Job")

    monkeypatch.setattr(manual_extraction, "fetch_page", fake_fetch)
    resp = client.post("/api/v1/jobs/manual", json={"url": kwargs.get("url", "https://example.com/jobs/1")})
    assert resp.status_code == 201, resp.text
    return resp.json()


@pytest.fixture()
def docs_dir(tmp_path, monkeypatch):
    """Points generated documents (and the master resume folder beside it) at a temp dir."""
    monkeypatch.setattr(settings, "GENERATED_DOCUMENTS_DIR", str(tmp_path / "generated_documents"))
    return tmp_path


_PROFILE = {
    "extraction": {
        "contact": {"name": "Jane Doe", "email": "jane@example.com"},
        "employment": [],
        "education": [],
        "skills": [],
        "certifications": [],
        "projects": [],
    },
    "approved_claims": [],
    "resume_filename": "resume.docx",
}


# --- structured data extraction (pure functions) ---------------------------


def test_structured_data_extracts_a_job_posting():
    posting = manual_extraction.extract_from_structured_data([_jsonld()], "https://example.com/jobs/1")
    assert posting is not None
    assert posting.title == "Data Analyst"
    assert posting.company == "Globex"
    assert posting.location == "Boston, MA"


def test_structured_data_requires_title_and_company():
    no_company = json.dumps({"@type": "JobPosting", "title": "Analyst"})
    no_title = json.dumps({"@type": "JobPosting", "hiringOrganization": {"name": "Globex"}})
    assert manual_extraction.extract_from_structured_data([no_company], "u") is None
    assert manual_extraction.extract_from_structured_data([no_title], "u") is None


def test_structured_data_handles_graph_and_ignores_other_types():
    graph = json.dumps({"@graph": [{"@type": "WebSite"}, json.loads(_jsonld())]})
    other = json.dumps({"@type": "Article", "title": "News", "hiringOrganization": {"name": "X"}})
    assert manual_extraction.extract_from_structured_data([graph], "u") is not None
    assert manual_extraction.extract_from_structured_data([other, "not json"], "u") is None


class _FakeExtractionRouter:
    def __init__(self, extraction: JobPostingExtraction):
        self._extraction = extraction

    def __call__(self, db):
        return self

    async def get_structured(self, **kwargs):
        return self._extraction


async def test_llm_fallback_refuses_pages_that_are_not_job_postings(monkeypatch):
    fake = _FakeExtractionRouter(JobPostingExtraction(is_job_posting=False))
    monkeypatch.setattr(manual_extraction, "ModelRouter", fake)
    with pytest.raises(ManualExtractionError):
        await manual_extraction.extract_job_posting(
            None, url="https://example.com/news", json_ld_blocks=[], body_text="Some article text", page_title="News"
        )


async def test_llm_fallback_never_accepts_a_posting_without_title_and_company(monkeypatch):
    fake = _FakeExtractionRouter(JobPostingExtraction(is_job_posting=True, title="Analyst", company=None))
    monkeypatch.setattr(manual_extraction, "ModelRouter", fake)
    with pytest.raises(ManualExtractionError):
        await manual_extraction.extract_job_posting(
            None, url="https://example.com/x", json_ld_blocks=[], body_text="text", page_title=None
        )


async def test_llm_fallback_accepts_a_real_posting(monkeypatch):
    fake = _FakeExtractionRouter(
        JobPostingExtraction(is_job_posting=True, title="Analyst", company="Globex", description="Do things")
    )
    monkeypatch.setattr(manual_extraction, "ModelRouter", fake)
    posting = await manual_extraction.extract_job_posting(
        None, url="https://example.com/x", json_ld_blocks=[], body_text="text", page_title=None
    )
    assert (posting.title, posting.company) == ("Analyst", "Globex")


# --- adding, listing, duplicates ---------------------------------------------


def test_add_job_by_link_and_list(app_and_db, monkeypatch):
    client, _ = app_and_db
    job = _add_job(client, monkeypatch)
    assert job["title"] == "Data Analyst"
    assert job["company"] == "Globex"
    assert job["application_status"] == "not_started"
    assert job["evaluation"] is None
    assert job["already_existed"] is False

    listed = client.get("/api/v1/jobs").json()
    assert [j["id"] for j in listed] == [job["id"]]


def test_adding_the_same_job_twice_says_so_and_keeps_one_row(app_and_db, monkeypatch):
    client, _ = app_and_db
    first = _add_job(client, monkeypatch)
    second = _add_job(client, monkeypatch)
    assert second["id"] == first["id"]
    assert second["already_existed"] is True
    assert len(client.get("/api/v1/jobs").json()) == 1


def test_page_that_is_not_a_job_is_rejected_cleanly(app_and_db, monkeypatch):
    client, _ = app_and_db

    async def fake_fetch(url: str) -> FetchedPage:
        return FetchedPage(url=url, json_ld_blocks=[], body_text="An article about weather.", title="News")

    monkeypatch.setattr(manual_extraction, "fetch_page", fake_fetch)
    resp = client.post("/api/v1/jobs/manual", json={"url": "https://example.com/news"})
    assert resp.status_code == 422
    assert client.get("/api/v1/jobs").json() == []


# --- scoring -----------------------------------------------------------------


class _FakeMatchRouter:
    def __init__(self, db):
        pass

    async def get_structured(self, **kwargs):
        return ResumeJobMatchResult(
            match_percentage=72, summary="Good SQL fit, short on years.", gaps=["3+ years of experience"]
        )


def test_scoring_stores_a_simple_result(app_and_db, monkeypatch):
    client, _ = app_and_db
    job = _add_job(client, monkeypatch)
    monkeypatch.setattr("app.services.qualification.pipeline.ModelRouter", _FakeMatchRouter)

    resp = client.post(f"/api/v1/jobs/{job['id']}/qualify")
    assert resp.status_code == 200
    evaluation = resp.json()["evaluation"]
    assert evaluation["overall_score"] == 0.72
    assert evaluation["summary"] == "Good SQL fit, short on years."
    assert evaluation["gaps"] == ["3+ years of experience"]

    again = client.post(f"/api/v1/jobs/{job['id']}/qualify")
    assert again.status_code == 200
    assert again.json()["evaluation"]["overall_score"] == 0.72  # re-scoring updates, not duplicates


def test_scoring_without_an_llm_is_a_clear_error_not_a_crash(app_and_db, monkeypatch):
    client, _ = app_and_db
    job = _add_job(client, monkeypatch)
    resp = client.post(f"/api/v1/jobs/{job['id']}/qualify")
    assert resp.status_code == 502
    assert client.get("/api/v1/jobs").json()[0]["evaluation"] is None


def test_adding_a_job_never_scores_it(app_and_db, monkeypatch):
    client, _ = app_and_db
    called = []

    class _Boom(_FakeMatchRouter):
        async def get_structured(self, **kwargs):
            called.append(1)
            return await super().get_structured(**kwargs)

    monkeypatch.setattr("app.services.qualification.pipeline.ModelRouter", _Boom)
    _add_job(client, monkeypatch)
    assert called == []


# --- proceed / undo / delete -------------------------------------------------


def test_proceed_and_undo(app_and_db, monkeypatch):
    client, _ = app_and_db
    job = _add_job(client, monkeypatch)
    assert client.post(f"/api/v1/jobs/{job['id']}/proceed").json()["application_status"] == "proceeding"
    assert client.post(f"/api/v1/jobs/{job['id']}/unproceed").json()["application_status"] == "not_started"
    assert client.post("/api/v1/jobs/9999/unproceed").status_code == 404


def test_delete_job_removes_it_and_its_generated_files(app_and_db, monkeypatch, docs_dir):
    client, SessionLocal = app_and_db
    job = _add_job(client, monkeypatch)

    inside = docs_dir / "generated_documents" / f"job_{job['id']}" / "resume.docx"
    inside.parent.mkdir(parents=True)
    inside.write_bytes(b"x")
    with SessionLocal() as db:
        db.add(
            GeneratedDocument(
                job_id=job["id"],
                document_type="resume",
                local_path=str(inside),
                format="docx",
                generated_at=dt.datetime.now(dt.timezone.utc),
                template_version="t",
                source_claim_ids=[],
                content_hash="h",
            )
        )
        db.commit()

    assert client.delete(f"/api/v1/jobs/{job['id']}").status_code == 204
    assert client.get(f"/api/v1/jobs/{job['id']}").status_code == 404
    assert not inside.exists()
    assert client.delete(f"/api/v1/jobs/{job['id']}").status_code == 404


def test_delete_never_removes_files_outside_the_documents_folder(app_and_db, monkeypatch, docs_dir):
    client, SessionLocal = app_and_db
    job = _add_job(client, monkeypatch)
    outside = docs_dir / "important.txt"
    outside.write_text("keep me")
    with SessionLocal() as db:
        db.add(
            GeneratedDocument(
                job_id=job["id"],
                document_type="resume",
                local_path=str(outside),
                format="txt",
                generated_at=dt.datetime.now(dt.timezone.utc),
                template_version="t",
                source_claim_ids=[],
                content_hash="h",
            )
        )
        db.commit()

    assert client.delete(f"/api/v1/jobs/{job['id']}").status_code == 204
    assert outside.exists()


# --- tailoring ---------------------------------------------------------------

_BULLETS = ["Reduced report time by 40% using SQL", "Led a team of 5 analysts", "Built Tableau dashboards"]


def _write_master(docs_dir) -> bytes:
    doc = Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("Analyst, Acme | 2020 – Present")
    for text in _BULLETS:
        doc.add_paragraph(text, style="List Bullet")
    master_dir = docs_dir / "master_resume"
    master_dir.mkdir(parents=True, exist_ok=True)
    path = master_dir / "master.docx"
    doc.save(str(path))
    return path.read_bytes()


class _FakePlanRouter:
    """Returns the plan it was built with for every call."""

    plan: TailorPlan

    def __init__(self, db):
        pass

    async def get_structured(self, **kwargs):
        return type(self).plan


def _plan(order: list[int], changelog=None) -> TailorPlan:
    return TailorPlan(
        blocks=[PlannedBlock(block_id=0, bullets=[PlannedBullet(bullet_id=i, text=_BULLETS[i]) for i in order])],
        changelog=changelog or [],
    )


def _docx_bullets(content: bytes) -> list[str]:
    return [p.text for p in Document(io.BytesIO(content)).paragraphs if p.style.name == "List Bullet"]


def test_tailoring_requires_proceed_first(app_and_db, monkeypatch, docs_dir):
    client, _ = app_and_db
    job = _add_job(client, monkeypatch)
    assert client.post(f"/api/v1/jobs/{job['id']}/tailor").status_code == 409


def test_tailoring_without_a_master_explains_what_to_do(app_and_db, monkeypatch, docs_dir):
    client, _ = app_and_db
    job = _add_job(client, monkeypatch)
    client.post(f"/api/v1/jobs/{job['id']}/proceed")
    resp = client.post(f"/api/v1/jobs/{job['id']}/tailor")
    assert resp.status_code == 409
    assert "Word" in resp.json()["detail"]


def test_tailoring_reorders_a_copy_and_leaves_the_master_untouched(app_and_db, monkeypatch, docs_dir):
    client, _ = app_and_db
    client.post("/api/v1/profile/commit", json=_PROFILE)
    job = _add_job(client, monkeypatch)
    client.post(f"/api/v1/jobs/{job['id']}/proceed")
    master_before = _write_master(docs_dir)

    _FakePlanRouter.plan = _plan([2, 0, 1], changelog=["Gap: the job wants Kubernetes; not in your resume."])
    monkeypatch.setattr("app.services.resume.generation.ModelRouter", _FakePlanRouter)

    resp = client.post(f"/api/v1/jobs/{job['id']}/tailor")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["used_original_wording"] is False
    assert body["changelog"] == [
        "Reordered the bullets in 1 role so the ones closest to this job come first.",
        "Gap: the job wants Kubernetes; not in your resume.",
    ]
    assert sorted(d["document_type"] for d in body["documents"]) == ["changelog", "resume"]

    resume_doc = next(d for d in body["documents"] if d["document_type"] == "resume")
    download = client.get(f"/api/v1/jobs/documents/{resume_doc['id']}/download")
    assert download.status_code == 200
    # The header is percent-encoded (spaces become %20); browsers decode it back.
    assert "Jane Doe_Resume_Globex_" in unquote(download.headers["content-disposition"])
    assert _docx_bullets(download.content) == [_BULLETS[2], _BULLETS[0], _BULLETS[1]]
    heading_line = Document(io.BytesIO(download.content)).paragraphs[1].text
    assert "\u2013" not in heading_line and "2020 - Present" in heading_line  # en dash became a hyphen

    assert (docs_dir / "master_resume" / "master.docx").read_bytes() == master_before

    listed = client.get("/api/v1/jobs").json()[0]
    assert sorted(d["document_type"] for d in listed["documents"]) == ["changelog", "resume"]


class _SequenceRouter:
    """Returns each plan in turn, and remembers what it was asked."""

    plans: list[TailorPlan] = []
    calls: list[list] = []

    def __init__(self, db):
        pass

    async def get_structured(self, **kwargs):
        type(self).calls.append(kwargs["messages"])
        plans = type(self).plans
        return plans[min(len(type(self).calls) - 1, len(plans) - 1)]


def test_a_plan_that_changes_nothing_is_sent_back_once_and_false_claims_are_dropped(app_and_db, monkeypatch, docs_dir):
    client, _ = app_and_db
    client.post("/api/v1/profile/commit", json=_PROFILE)
    job = _add_job(client, monkeypatch)
    client.post(f"/api/v1/jobs/{job['id']}/proceed")
    _write_master(docs_dir)

    claims = ["Reordered bullets to prioritize SQL.", "Rephrased bullets toward the job.", "Master lacks Kubernetes."]
    _SequenceRouter.calls = []
    _SequenceRouter.plans = [_plan([0, 1, 2], changelog=claims), _plan([1, 2, 0], changelog=claims)]
    monkeypatch.setattr("app.services.resume.generation.ModelRouter", _SequenceRouter)

    body = client.post(f"/api/v1/jobs/{job['id']}/tailor").json()
    assert len(_SequenceRouter.calls) == 2  # asked again once, with the reason
    assert "identical to the master" in _SequenceRouter.calls[1][-1]["content"]
    assert body["changelog"] == [
        "Reordered the bullets in 1 role so the ones closest to this job come first.",
        "Master lacks Kubernetes.",
    ]
    resume_doc = next(d for d in body["documents"] if d["document_type"] == "resume")
    download = client.get(f"/api/v1/jobs/documents/{resume_doc['id']}/download")
    assert _docx_bullets(download.content) == [_BULLETS[1], _BULLETS[2], _BULLETS[0]]


def test_a_plan_that_still_changes_nothing_says_so_and_never_claims_a_change(app_and_db, monkeypatch, docs_dir):
    client, _ = app_and_db
    client.post("/api/v1/profile/commit", json=_PROFILE)
    job = _add_job(client, monkeypatch)
    client.post(f"/api/v1/jobs/{job['id']}/proceed")
    _write_master(docs_dir)

    _SequenceRouter.calls = []
    _SequenceRouter.plans = [_plan([0, 1, 2], changelog=["Reordered bullets to prioritize SQL."])]
    monkeypatch.setattr("app.services.resume.generation.ModelRouter", _SequenceRouter)

    body = client.post(f"/api/v1/jobs/{job['id']}/tailor").json()
    assert len(_SequenceRouter.calls) == 2
    assert len(body["changelog"]) == 1 and body["changelog"][0].startswith("Nothing was changed")


def test_tailoring_falls_back_to_the_master_when_the_ai_plan_fails_checks(app_and_db, monkeypatch, docs_dir):
    client, _ = app_and_db
    client.post("/api/v1/profile/commit", json=_PROFILE)
    job = _add_job(client, monkeypatch)
    client.post(f"/api/v1/jobs/{job['id']}/proceed")
    _write_master(docs_dir)

    _FakePlanRouter.plan = _plan([0, 1])  # drops a bullet: must be rejected
    monkeypatch.setattr("app.services.resume.generation.ModelRouter", _FakePlanRouter)

    resp = client.post(f"/api/v1/jobs/{job['id']}/tailor")
    assert resp.status_code == 200
    body = resp.json()
    assert body["used_original_wording"] is True
    assert any("accuracy checks" in note for note in body["changelog"])

    resume_doc = next(d for d in body["documents"] if d["document_type"] == "resume")
    download = client.get(f"/api/v1/jobs/documents/{resume_doc['id']}/download")
    assert _docx_bullets(download.content) == _BULLETS


def test_recreating_a_tailored_resume_replaces_the_old_one(app_and_db, monkeypatch, docs_dir):
    client, _ = app_and_db
    client.post("/api/v1/profile/commit", json=_PROFILE)
    job = _add_job(client, monkeypatch)
    client.post(f"/api/v1/jobs/{job['id']}/proceed")
    _write_master(docs_dir)
    _FakePlanRouter.plan = _plan([1, 0, 2])
    monkeypatch.setattr("app.services.resume.generation.ModelRouter", _FakePlanRouter)

    client.post(f"/api/v1/jobs/{job['id']}/tailor")
    second = client.post(f"/api/v1/jobs/{job['id']}/tailor").json()

    resume_doc = next(d for d in second["documents"] if d["document_type"] == "resume")
    download = client.get(f"/api/v1/jobs/documents/{resume_doc['id']}/download")
    assert download.status_code == 200  # the new file still exists after the old one was cleared
    assert len(client.get("/api/v1/jobs").json()[0]["documents"]) == 2


def test_download_refuses_unknown_and_out_of_folder_documents(app_and_db, monkeypatch, docs_dir):
    client, SessionLocal = app_and_db
    job = _add_job(client, monkeypatch)
    secret = docs_dir / "secret.txt"
    secret.write_text("private")
    with SessionLocal() as db:
        doc = GeneratedDocument(
            job_id=job["id"],
            document_type="resume",
            local_path=str(secret),
            format="txt",
            generated_at=dt.datetime.now(dt.timezone.utc),
            template_version="t",
            source_claim_ids=[],
            content_hash="h",
        )
        db.add(doc)
        db.commit()
        doc_id = doc.id

    assert client.get(f"/api/v1/jobs/documents/{doc_id}/download").status_code == 404
    assert client.get("/api/v1/jobs/documents/9999/download").status_code == 404


# --- work eligibility answers -------------------------------------------------


def test_work_eligibility_answers_need_a_profile_then_round_trip(app_and_db):
    client, _ = app_and_db
    assert client.get("/api/v1/profile/answers").status_code == 404

    client.post("/api/v1/profile/commit", json=_PROFILE)
    saves = {
        "work_authorization": {"value_type": "str", "value": "US citizen"},
        "sponsorship_required": {"value_type": "bool", "value": False},
        "security_clearance": {"value_type": "str", "value": "None"},
    }
    for key, payload in saves.items():
        assert client.put(f"/api/v1/profile/answers/{key}", json=payload).status_code == 200

    answers = {a["answer_key"]: a["value"] for a in client.get("/api/v1/profile/answers").json()}
    assert answers == {
        "work_authorization": "US citizen",
        "sponsorship_required": False,
        "security_clearance": "None",
    }

    client.put(
        "/api/v1/profile/answers/sponsorship_required", json={"value_type": "bool", "value": True}
    )
    updated = {a["answer_key"]: a["value"] for a in client.get("/api/v1/profile/answers").json()}
    assert updated["sponsorship_required"] is True
    assert len(client.get("/api/v1/profile/answers").json()) == 3  # updated in place, not duplicated


def test_original_resume_needs_proceed_first(app_and_db, monkeypatch, docs_dir):
    client, _ = app_and_db
    job = _add_job(client, monkeypatch)
    assert client.post(f"/api/v1/jobs/{job['id']}/original-resume").status_code == 409


def test_original_resume_without_a_master_explains_what_to_do(app_and_db, monkeypatch, docs_dir):
    client, _ = app_and_db
    client.post("/api/v1/profile/commit", json=_PROFILE)
    job = _add_job(client, monkeypatch)
    client.post(f"/api/v1/jobs/{job['id']}/proceed")
    resp = client.post(f"/api/v1/jobs/{job['id']}/original-resume")
    assert resp.status_code == 409 and "Upload your resume" in resp.json()["detail"]


def test_original_resume_is_an_unchanged_copy_and_the_master_is_untouched(app_and_db, monkeypatch, docs_dir):
    client, _ = app_and_db
    client.post("/api/v1/profile/commit", json=_PROFILE)
    job = _add_job(client, monkeypatch)
    client.post(f"/api/v1/jobs/{job['id']}/proceed")
    master_before = _write_master(docs_dir)

    resp = client.post(f"/api/v1/jobs/{job['id']}/original-resume")
    assert resp.status_code == 200, resp.text
    docs = resp.json()
    assert [d["document_type"] for d in docs] == ["resume"] and docs[0]["template_version"] == "original-1"
    download = client.get(f"/api/v1/jobs/documents/{docs[0]['id']}/download")
    assert "Jane Doe_Resume_Globex_" in unquote(download.headers["content-disposition"])
    assert download.content == master_before
    assert (docs_dir / "master_resume" / "master.docx").read_bytes() == master_before

    listed = client.get("/api/v1/jobs").json()[0]
    assert [d["document_type"] for d in listed["documents"]] == ["resume"]


def test_mark_applied_with_a_date_and_take_it_back(app_and_db, monkeypatch):
    client, _ = app_and_db
    job = _add_job(client, monkeypatch)
    assert job["applied_at"] is None
    out = client.post(f"/api/v1/jobs/{job['id']}/applied", json={"applied_on": "2026-10-01"}).json()
    assert out["applied_at"].startswith("2026-10-01") and out["applied_via"] == "manual"
    out = client.post(f"/api/v1/jobs/{job['id']}/applied").json()  # no date: today
    assert out["applied_at"] is not None
    assert client.post(f"/api/v1/jobs/{job['id']}/applied", json={"applied_on": "2999-01-01"}).status_code == 422
    out = client.post(f"/api/v1/jobs/{job['id']}/unapplied").json()
    assert out["applied_at"] is None and out["applied_via"] is None
    assert client.post("/api/v1/jobs/9999/applied").status_code == 404


class _FakeDraftRouter:
    def __init__(self, db):
        pass

    async def get_structured(self, **kwargs):
        from app.services.followup.draft import FollowUpDraft

        user = kwargs["messages"][-1]["content"]
        return FollowUpDraft(subject="Checking in — Data Analyst", body=f"Hello, {user[:20]} – thanks.")


def test_follow_up_draft_is_returned_not_saved_and_has_no_long_dashes(app_and_db, monkeypatch):
    client, _ = app_and_db
    job = _add_job(client, monkeypatch)
    monkeypatch.setattr("app.services.followup.draft.ModelRouter", _FakeDraftRouter)
    resp = client.post(f"/api/v1/jobs/{job['id']}/follow-up-draft", json={"kind": "after_interview"})
    assert resp.status_code == 200
    out = resp.json()
    assert out["subject"] == "Checking in , Data Analyst"
    assert "—" not in out["body"] and "–" not in out["body"]
    assert client.post(f"/api/v1/jobs/{job['id']}/follow-up-draft").status_code == 200  # kind is optional
    assert client.post("/api/v1/jobs/9999/follow-up-draft").status_code == 404


def test_follow_up_draft_without_an_llm_is_a_clear_error(app_and_db, monkeypatch):
    client, _ = app_and_db
    job = _add_job(client, monkeypatch)
    assert client.post(f"/api/v1/jobs/{job['id']}/follow-up-draft").status_code == 502


def test_interviewing_implies_applied_and_follow_up_can_be_cleared(app_and_db, monkeypatch):
    client, _ = app_and_db
    job = _add_job(client, monkeypatch)
    out = client.post(f"/api/v1/jobs/{job['id']}/interviewing").json()
    assert out["interviewing_at"] and out["applied_at"]
    done = client.post(f"/api/v1/jobs/{job['id']}/follow-up-done").json()
    assert done["followup_done_at"]
    out = client.post(f"/api/v1/jobs/{job['id']}/not-interviewing").json()
    assert out["interviewing_at"] is None and out["applied_at"]
    out = client.post(f"/api/v1/jobs/{job['id']}/interviewing").json()
    out = client.post(f"/api/v1/jobs/{job['id']}/unapplied").json()
    assert out["applied_at"] is None and out["interviewing_at"] is None
    assert client.post("/api/v1/jobs/9999/interviewing").status_code == 404


def test_a_busy_ai_provider_gets_a_friendly_message_with_no_technical_detail(app_and_db, monkeypatch):
    import asyncio

    import httpx

    from app.services.llm import router as llm_router
    from app.services.llm.exceptions import StructuredOutputError
    from app.services.qualification.pipeline import QualificationError, score_against

    monkeypatch.setattr(llm_router, "BUSY_BACKOFF_SECONDS", 0)
    monkeypatch.setattr(llm_router.settings, "PRIMARY_FAST_MODEL", "m")
    request = httpx.Request("POST", "https://openrouter.ai/api/v1/chat/completions")

    class Busy:
        async def chat_completion(self, **kw):
            raise httpx.HTTPStatusError("429 Too Many Requests", request=request, response=httpx.Response(429, request=request))

    _, SessionLocal = app_and_db
    with SessionLocal() as db:
        router = llm_router.ModelRouter(db, client=Busy())
        try:
            asyncio.run(router.get_structured(purpose="t", prompt_version="v1", messages=[], response_model=__import__("app.services.llm.schemas", fromlist=["ResumeJobMatchResult"]).ResumeJobMatchResult))
            raise AssertionError("should have failed")
        except StructuredOutputError as exc:
            assert exc.busy is True

        monkeypatch.setattr("app.services.qualification.pipeline.ModelRouter", lambda db: llm_router.ModelRouter(db, client=Busy()))
        try:
            asyncio.run(score_against(db, title="A", company="B", description="C"))
            raise AssertionError("should have failed")
        except QualificationError as exc:
            assert str(exc) == "Scoring is busy right now. Try again in a minute."


def test_skills_and_summary_numbered_wrongly_by_the_model_are_matched_by_what_they_list():
    from app.services.llm.schemas import PlannedSkills, PlannedSummary
    from app.services.resume.docx_editor import TextPart
    from app.services.resume.generation import repair_part_ids

    parts = [
        TextPart(100, "summary", None, "A long enough summary paragraph for a revenue cycle consultant.", 0),
        TextPart(101, "skills", None, "Power BI, Excel", 0, label="Data: ", items=["Power BI", "Excel"], seps=[", "]),
        TextPart(102, "skills", None, "Epic, SSI", 0, label="Systems: ", items=["Epic", "SSI"], seps=[", "]),
    ]
    plan = TailorPlan(
        summaries=[PlannedSummary(part_id=0, text="Reworded.")],
        skills=[PlannedSkills(part_id=1, items=["SSI", "Epic"]), PlannedSkills(part_id=2, items=["Excel", "Power BI", "Change Management"])],
    )
    fixed = repair_part_ids(plan, parts)
    assert fixed.summaries[0].part_id == 100
    assert [(s.part_id, s.items[0]) for s in fixed.skills] == [(102, "SSI"), (101, "Excel")]


def test_skills_lines_written_as_bullets_are_edited_as_skills_not_as_a_role():
    from app.services.resume import docx_editor

    doc = Document()
    doc.add_paragraph("Professional Experience")
    doc.add_paragraph("Analyst, Acme | 2020 - Present")
    doc.add_paragraph("Reduced report time by 40% using SQL", style="List Bullet")
    doc.add_paragraph("Led a team of 5 analysts", style="List Bullet")
    doc.add_paragraph("Skills and Proficiencies")
    doc.add_paragraph("Data Analytics: Power BI, Advanced Excel", style="List Bullet")
    doc.add_paragraph("Healthcare Strategy: EHR Implementation, Patient Financial Services (A/R, Billing)", style="List Bullet")

    blocks = docx_editor.find_blocks(doc)
    parts = docx_editor.find_text_parts(doc)
    assert [len(b.bullets) for b in blocks] == [2]  # only the role's bullets
    skills = [p for p in parts if p.kind == "skills"]
    assert [p.label for p in skills] == ["Data Analytics: ", "Healthcare Strategy: "]
    assert skills[1].items == ["EHR Implementation", "Patient Financial Services (A/R, Billing)"]

    docx_editor.apply_text_plan(parts, {}, {skills[0].id: ["Advanced Excel", "Power BI"]})
    assert "Data Analytics: Advanced Excel, Power BI" in [p.text for p in doc.paragraphs]


def test_notes_name_new_keywords_and_where_they_were_added():
    from app.services.llm.schemas import PlannedSummary
    from app.services.resume.docx_editor import Bullet, Block, TextPart
    from app.services.resume.generation import _change_notes

    master = "Engineered executive Power BI dashboards and briefed senior leadership. Led teams of consultants."
    blocks = [Block(0, "Senior Consultant | Tria Federal | Arlington, VA", [
        Bullet(0, None, "Engineered executive Power BI dashboards", ""),
        Bullet(1, None, "Led teams of consultants", ""),
    ])]
    parts = [TextPart(100, "summary", None, "Experienced consultant who builds Power BI dashboards.", 0)]
    plan = TailorPlan(
        summaries=[PlannedSummary(part_id=100, text="Experienced consultant who builds Power BI dashboards and predictive models.")],
        blocks=[PlannedBlock(block_id=0, bullets=[
            PlannedBullet(bullet_id=1, text="Led teams of consultants on stakeholder engagement"),
            PlannedBullet(bullet_id=0, text="Engineered executive Power BI dashboards"),
        ])],
    )
    notes = _change_notes(plan, blocks, parts, None, master)
    assert "New wording in your summary that your master doesn't use: predictive, models. Check each one is true for you." in notes
    assert any(n.startswith("New wording in bullet 1 of Senior Consultant | Tria Federal: stakeholder, engagement.") for n in notes)
    assert not any("dashboards" in n and n.startswith("New wording") for n in notes)  # a word the master already uses is never flagged


def test_tailoring_refuses_when_the_job_has_no_description(app_and_db, monkeypatch, docs_dir):
    client, _ = app_and_db
    job = _add_job(client, monkeypatch, description="<p>Short summary.</p>")
    client.post(f"/api/v1/jobs/{job['id']}/proceed")
    resp = client.post(f"/api/v1/jobs/{job['id']}/tailor")
    assert resp.status_code == 428
    assert "extension" in resp.json()["detail"]


def test_capturing_the_page_fills_in_the_job_saved_earlier_without_a_description(app_and_db, monkeypatch):
    client, _ = app_and_db
    first = _add_job(client, monkeypatch, description="<p>Short summary.</p>", url="https://search.example/abc")
    secret = client.post("/api/v1/system/pairing-secret").json()["pairing_secret"]
    token = client.post("/api/v1/system/pair", json={"pairing_secret": secret, "extension_origin": "chrome-extension://t"}).json()["extension_token"]
    body = {"url": "https://careers.globex.example/jobs/99", "page_title": "Data Analyst",
            "json_ld": [_jsonld(url="https://careers.globex.example/jobs/99")], "body_text": ""}
    resp = client.post("/api/v1/jobs/capture", json=body, headers={"X-Extension-Token": token})
    assert resp.status_code == 201, resp.text
    assert resp.json()["id"] == first["id"]  # the same job, not a copy
    assert resp.json()["already_existed"] is True
    assert len(client.get("/api/v1/jobs").json()) == 1
    detail = client.get(f"/api/v1/jobs/{first['id']}").json()
    assert "build reports" in detail["description"]


def test_capture_with_a_job_id_fills_in_that_exact_job(app_and_db, monkeypatch):
    client, _ = app_and_db
    first = _add_job(client, monkeypatch, description="<p>Short summary.</p>", url="https://search.example/abc")
    other = _add_job(client, monkeypatch, title="Data Engineer", url="https://search.example/other", description="<p>Short.</p>")
    secret = client.post("/api/v1/system/pairing-secret").json()["pairing_secret"]
    token = client.post("/api/v1/system/pair", json={"pairing_secret": secret, "extension_origin": "chrome-extension://t"}).json()["extension_token"]
    headers = {"X-Extension-Token": token}
    page = {"url": "https://careers.example/anything", "page_title": "Whatever", "json_ld": [_jsonld(title="Some Other Title", company="Other Co")], "body_text": ""}
    resp = client.post("/api/v1/jobs/capture", json={**page, "job_id": first["id"]}, headers=headers)
    assert resp.status_code == 201, resp.text
    assert resp.json()["id"] == first["id"] and resp.json()["title"] == "Data Analyst"  # the saved job keeps its own title
    assert "build reports" in client.get(f"/api/v1/jobs/{first['id']}").json()["description"]
    assert "build reports" not in client.get(f"/api/v1/jobs/{other['id']}").json()["description"]
    assert len(client.get("/api/v1/jobs").json()) == 2
    thin = {**page, "json_ld": [_jsonld(description="<p>Too short.</p>")]}
    assert client.post("/api/v1/jobs/capture", json={**thin, "job_id": other["id"]}, headers=headers).status_code == 422
