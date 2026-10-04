"""Profile editing (per-cell edits), master bullets view, and the dashboard refresh fallback."""

from __future__ import annotations

import docx
from fastapi.testclient import TestClient

_COMMIT_PAYLOAD = {
    "extraction": {
        "contact": {"name": "Jane Doe", "email": "jane@example.com"},
        "employment": [
            {
                "employer": "Acme Corp",
                "title": "Senior Data Analyst",
                "start_date": "2021-04",
                "end_date": None,
                "source_text": "Built Power BI dashboards using SQL.",
            }
        ],
        "education": [],
        "skills": ["SQL", "Power BI"],
        "certifications": [],
        "projects": [],
    },
    "approved_claims": [],
    "resume_filename": "resume.docx",
}


def _commit(client):
    return client.post("/api/v1/profile/commit", json=_COMMIT_PAYLOAD).json()


def test_profile_ids_and_employment_edit(app_and_db):
    client, _ = app_and_db
    profile = _commit(client)
    role = profile["employment_history"][0]
    assert "id" in role

    resp = client.patch(
        f"/api/v1/profile/employment/{role['id']}",
        json={"title": "Lead Analyst", "location": "Boston, MA", "end_date": "2024-06"},
    )
    assert resp.status_code == 200
    updated = resp.json()["employment_history"][0]
    assert updated["title"] == "Lead Analyst"
    assert updated["location"] == "Boston, MA"
    assert updated["end_date"] == "2024-06-01"
    assert updated["employer"] == "Acme Corp"  # untouched

    # Blank required fields are ignored, never saved as empty.
    resp = client.patch(f"/api/v1/profile/employment/{role['id']}", json={"employer": "  "})
    assert resp.json()["employment_history"][0]["employer"] == "Acme Corp"

    # Clearing the end date means current role.
    resp = client.patch(f"/api/v1/profile/employment/{role['id']}", json={"end_date": ""})
    assert resp.json()["employment_history"][0]["end_date"] is None


def test_employment_add_delete_and_unknown_id(app_and_db):
    client, _ = app_and_db
    _commit(client)
    added = client.post("/api/v1/profile/employment", json={"employer": "Globex", "title": "Analyst"})
    assert added.status_code == 201
    assert len(added.json()["employment_history"]) == 2
    new_id = added.json()["employment_history"][-1]["id"]
    after = client.delete(f"/api/v1/profile/employment/{new_id}")
    assert len(after.json()["employment_history"]) == 1
    assert client.patch("/api/v1/profile/employment/9999", json={"title": "x"}).status_code == 404


def test_education_and_skills_edit(app_and_db):
    client, _ = app_and_db
    _commit(client)
    added = client.post("/api/v1/profile/education", json={"institution": "State U", "degree": "BS"}).json()
    edu = added["education"][0]
    updated = client.patch(
        f"/api/v1/profile/education/{edu['id']}", json={"field": "Economics", "end_date": "2019"}
    ).json()["education"][0]
    assert updated["field"] == "Economics"
    assert updated["end_date"] == "2019-01-01"
    assert client.delete(f"/api/v1/profile/education/{edu['id']}").json()["education"] == []

    with_skill = client.post("/api/v1/profile/skills", json={"canonical_skill": " Tableau "}).json()
    names = [s["canonical_skill"] for s in with_skill["skills"]]
    assert "Tableau" in names
    again = client.post("/api/v1/profile/skills", json={"canonical_skill": "tableau"}).json()
    assert [s["canonical_skill"] for s in again["skills"]].count("Tableau") == 1  # no duplicates
    sid = next(s["id"] for s in again["skills"] if s["canonical_skill"] == "Tableau")
    left = client.delete(f"/api/v1/profile/skills/{sid}").json()["skills"]
    assert "Tableau" not in [s["canonical_skill"] for s in left]
    assert client.post("/api/v1/profile/skills", json={"canonical_skill": "  "}).status_code == 422


def test_patch_profile_ignores_blank_name_and_email(app_and_db):
    client, _ = app_and_db
    _commit(client)
    resp = client.patch("/api/v1/profile", json={"name": "", "email": " ", "phone": "555-0111"})
    body = resp.json()
    assert body["name"] == "Jane Doe"
    assert body["email"] == "jane@example.com"
    assert body["phone"] == "555-0111"


def test_edits_need_a_profile(app_and_db):
    client, _ = app_and_db
    assert client.post("/api/v1/profile/employment", json={"title": "x"}).status_code == 404
    assert client.post("/api/v1/profile/skills", json={"canonical_skill": "SQL"}).status_code == 404


def test_master_roles_reads_bullets_without_touching_the_file(app_and_db, tmp_path, monkeypatch):
    client, _ = app_and_db
    path = tmp_path / "master.docx"
    d = docx.Document()
    d.add_paragraph("Senior Analyst | Acme Corp")
    d.add_paragraph("Built dashboards for finance.", style="List Bullet")
    d.add_paragraph("Cut reporting time by 30%.", style="List Bullet")
    d.save(str(path))
    before = path.read_bytes()

    from app.services.resume import master_view

    monkeypatch.setattr(master_view, "master_path", lambda: path)
    roles = client.get("/api/v1/profile/master/roles").json()
    assert roles[0]["bullets"] == ["Built dashboards for finance.", "Cut reporting time by 30%."]
    assert "Acme Corp" in roles[0]["context"]
    assert path.read_bytes() == before


def test_dashboard_refresh_falls_back_to_index(tmp_path, monkeypatch):
    dist = tmp_path / "dist"
    (dist / "assets").mkdir(parents=True)
    (dist / "index.html").write_text("<html>app shell</html>")
    (dist / "assets" / "x.js").write_text("console.log(1)")
    (tmp_path / "secret.txt").write_text("nope")

    from app import main

    monkeypatch.setattr(main, "DASHBOARD_DIST", dist)
    client = TestClient(main.create_app(), base_url="http://127.0.0.1")  # the local-only guard rejects other hosts
    assert "app shell" in client.get("/app/jobs").text  # refresh on a tab address
    assert "app shell" in client.get("/app").text
    assert client.get("/app/assets/x.js").text == "console.log(1)"  # real files still served
    assert "nope" not in client.get("/app/../secret.txt").text
    assert client.get("/api/v1/nonexistent").status_code == 404


def test_degree_abbreviation_is_saved_and_can_be_changed(app_and_db):
    client, _ = app_and_db
    _commit(client)
    edu = client.post(
        "/api/v1/profile/education",
        json={"institution": "State U", "degree": "Bachelor's Degree", "degree_short": " B.S. "},
    ).json()["education"][0]
    assert edu["degree_short"] == "B.S."
    changed = client.patch(f"/api/v1/profile/education/{edu['id']}", json={"degree_short": "B.A."}).json()["education"][0]
    assert changed["degree_short"] == "B.A." and changed["degree"] == "Bachelor's Degree"
