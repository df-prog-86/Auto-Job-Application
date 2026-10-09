from docx import Document

from app.services.llm.schemas import PlannedSkills, PlannedSummary, TailorPlan
from app.services.resume import docx_editor as de
from app.services.resume.validation import skills_to_dict, summaries_to_dict, validate_plan

SUMMARY = "Revenue cycle consultant with ten years of experience leading denial reduction and billing process improvement for hospital systems."


def _doc():
    d = Document()
    d.add_paragraph("SUMMARY")
    d.add_paragraph(SUMMARY)
    d.add_paragraph("EXPERIENCE")
    d.add_paragraph("Consultant | Acme Health")
    d.add_paragraph("Reduced denials by 20%", style="List Bullet")
    d.add_paragraph("SKILLS")
    q = d.add_paragraph()
    q.add_run("Tools: ").bold = True
    q.add_run("Epic (Resolute, Cadence), SQL, Excel, Tableau")
    d.add_paragraph("EDUCATION")
    d.add_paragraph("UCF")
    return d


def test_summary_and_skills_are_found_and_edited_in_place():
    d = _doc()
    blocks, parts = de.find_blocks(d), de.find_text_parts(d)
    assert [p.kind for p in parts] == ["summary", "skills"]
    plan = TailorPlan(
        summaries=[PlannedSummary(part_id=100, text=SUMMARY.replace("leading", "driving"))],
        skills=[PlannedSkills(part_id=101, items=["tableau", "SQL", "Epic (Resolute, Cadence)", "Excel"])],
    )
    assert validate_plan(plan, blocks, de.full_text(d), False, parts) == []
    assert de.apply_text_plan(parts, summaries_to_dict(plan), skills_to_dict(plan, parts)) == 2
    texts = [p.text for p in d.paragraphs]
    assert "driving denial reduction" in texts[1]
    skills = d.paragraphs[6]
    assert skills.text == "Tools: Tableau, SQL, Epic (Resolute, Cadence), Excel"
    assert skills.runs[0].text == "Tools: " and skills.runs[0].bold


def test_skills_cannot_be_added_or_dropped():
    d = _doc()
    blocks, parts = de.find_blocks(d), de.find_text_parts(d)
    bad = TailorPlan(skills=[PlannedSkills(part_id=101, items=["Tableau", "SQL", "Python", "Excel"])])
    assert validate_plan(bad, blocks, de.full_text(d), False, parts)


def test_light_summary_cannot_invent_numbers_or_grow():
    d = _doc()
    blocks, parts = de.find_blocks(d), de.find_text_parts(d)
    plan = TailorPlan(summaries=[PlannedSummary(part_id=100, text=SUMMARY.replace("ten", "15"))])
    assert validate_plan(plan, blocks, de.full_text(d), False, parts)
    longer = TailorPlan(summaries=[PlannedSummary(part_id=100, text=SUMMARY + " " + SUMMARY[:60])])
    assert validate_plan(longer, blocks, de.full_text(d), False, parts)


def test_sanitize_keeps_good_edits_and_restores_bad_ones():
    from app.services.llm.schemas import PlannedBlock, PlannedBullet
    from app.services.resume.validation import sanitize_plan

    d = _doc()
    blocks, parts = de.find_blocks(d), de.find_text_parts(d)
    bullet = blocks[0].bullets[0]
    plan = TailorPlan(
        blocks=[PlannedBlock(block_id=0, bullets=[PlannedBullet(bullet_id=bullet.id, text="Cut denials by 90% using Lean Six Sigma")])],
        summaries=[PlannedSummary(part_id=100, text=SUMMARY.replace("leading", "driving"))],
        skills=[PlannedSkills(part_id=101, items=["Tableau", "SQL", "Python", "Excel"])],
    )
    assert validate_plan(plan, blocks, de.full_text(d), True, parts)
    cleaned, notes = sanitize_plan(plan, blocks, de.full_text(d), True, parts)
    assert validate_plan(cleaned, blocks, de.full_text(d), True, parts) == []
    assert cleaned.blocks[0].bullets[0].text == bullet.text
    assert len(cleaned.summaries) == 1 and cleaned.skills == []
    assert notes


def test_compound_terms_made_of_resume_words_are_allowed():
    from app.services.resume.validation import _check_rewrite

    vocab = {"built", "an", "ai", "chatbot", "powered", "by", "gpt", "for", "claims"}
    assert _check_rewrite(1, "Built an AI chatbot powered by GPT.", "Built an AI-powered chatbot for claims.", vocab, True) == []
    assert _check_rewrite(1, "Built an AI chatbot powered by GPT.", "Built a Salesforce-powered chatbot.", vocab, True)
