"""Milestone 5: tailoring edits a copy of the master Word file in place."""

from docx import Document

from app.services.llm.schemas import PlannedBlock, PlannedBullet, TailorPlan
from app.services.resume import docx_editor
from app.services.resume.naming import resume_filename
from app.services.resume.validation import plan_to_dict, validate_plan


def _master():
    doc = Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("Analyst, Acme | 2020 – Present")
    for text in ["Reduced report time by 40% using SQL", "Led a team of 5 analysts", "Built Tableau dashboards"]:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_paragraph("Intern, Beta | 2019")
    doc.add_paragraph("Cleaned data in Excel", style="List Bullet")
    doc.add_paragraph("")
    doc.add_paragraph("")
    return doc


def _plan(blocks, order, texts=None):
    texts = texts or {}
    bullets = {b.id: b for blk in blocks for b in blk.bullets}
    return TailorPlan(
        blocks=[
            PlannedBlock(
                block_id=0,
                bullets=[PlannedBullet(bullet_id=i, text=texts.get(i, bullets[i].text)) for i in order],
            )
        ]
    )


def test_blocks_found_per_role():
    blocks = docx_editor.find_blocks(_master())
    assert [len(b.bullets) for b in blocks] == [3, 1]
    assert "Acme" in blocks[0].context


def test_reorder_keeps_layout_and_changes_order_only():
    doc = _master()
    blocks = docx_editor.find_blocks(doc)
    plan = _plan(blocks, [2, 0, 1])
    assert validate_plan(plan, blocks, docx_editor.full_text(doc)) == []
    docx_editor.apply_plan(blocks, plan_to_dict(plan))
    texts = [p.text for p in doc.paragraphs]
    assert texts[2:5] == ["Built Tableau dashboards", "Reduced report time by 40% using SQL", "Led a team of 5 analysts"]
    assert all(p.style.name == "List Bullet" for p in doc.paragraphs[2:5])
    assert texts[1].startswith("Analyst, Acme") and texts[5].startswith("Intern, Beta")


def test_dropped_or_duplicated_bullets_rejected():
    doc = _master()
    blocks = docx_editor.find_blocks(doc)
    assert validate_plan(_plan(blocks, [0, 1]), blocks, docx_editor.full_text(doc))
    assert validate_plan(_plan(blocks, [0, 0, 1]), blocks, docx_editor.full_text(doc))


def test_invented_number_term_and_heavy_rewrite_rejected():
    doc = _master()
    blocks = docx_editor.find_blocks(doc)
    full = docx_editor.full_text(doc)
    assert validate_plan(_plan(blocks, [0, 1, 2], {0: "Reduced report time by 60% using SQL"}), blocks, full)
    assert validate_plan(_plan(blocks, [0, 1, 2], {2: "Built Tableau dashboards in Kubernetes"}), blocks, full)
    assert validate_plan(_plan(blocks, [0, 1, 2], {1: "Managed people"}), blocks, full)


def test_light_leadin_allowed_and_applied():
    doc = _master()
    blocks = docx_editor.find_blocks(doc)
    plan = _plan(blocks, [0, 1, 2], {0: "Streamlined reporting: reduced report time by 40% using SQL"})
    assert validate_plan(plan, blocks, docx_editor.full_text(doc)) == []
    assert docx_editor.apply_plan(blocks, plan_to_dict(plan)) == 1
    assert doc.paragraphs[2].text.startswith("Streamlined reporting")


def test_dash_cleanup_and_trailing_blank_strip():
    doc = _master()
    docx_editor.clean_dashes(doc)
    docx_editor.strip_trailing_blank_paragraphs(doc)
    assert "2020 - Present" in doc.paragraphs[1].text
    assert doc.paragraphs[-1].text == "Cleaned data in Excel"


def test_filename_rules():
    assert resume_filename("Jane Doe", "Mass General Brigham", "pdf", 2026) == "Jane Doe_Resume_Mass General Brigham_2026.pdf"
    assert resume_filename("Jane Doe", "A/B: Co", "docx", 2026) == "Jane Doe_Resume_AB Co_2026.docx"
